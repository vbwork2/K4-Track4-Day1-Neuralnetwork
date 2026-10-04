"""Strict JSON histories and lab-template workbook export using openpyxl."""
from __future__ import annotations
from copy import copy
import json
import math
from pathlib import Path
import re
import numpy as np
import openpyxl
from openpyxl.formula.translate import Translator
from openpyxl.workbook.properties import CalcProperties

FORMULA_COLUMNS = {"step0_gap_vs_lnC", "gap_val_minus_train", "delta_val_f1_vs_base", "beyond_noise"}
OPT_LABELS = {"sgd": "SGD", "sgd_momentum": "SGD+momentum", "adam": "Adam", "adamw": "AdamW"}


def json_safe(value):
    if isinstance(value, dict):
        return {str(k): json_safe(v) for k, v in value.items()}
    if isinstance(value, (list, tuple, np.ndarray)):
        return [json_safe(v) for v in value]
    if isinstance(value, np.generic):
        value = value.item()
    if isinstance(value, float) and not math.isfinite(value):
        return None
    if isinstance(value, Path):
        return str(value)
    return value


def write_json(path, value):
    target = Path(path)
    target.parent.mkdir(parents=True, exist_ok=True)
    temporary = target.with_suffix(target.suffix + ".tmp")
    temporary.write_text(json.dumps(json_safe(value), indent=2, allow_nan=False), encoding="utf-8")
    temporary.replace(target)


def save_result(result: dict, results_dir: str = "../results") -> str:
    exp_id = result["cfg"]["exp_id"]
    if not re.fullmatch(r"[A-Za-z0-9_.-]+", exp_id):
        raise ValueError("Experiment IDs must be safe file names.")
    target = Path(results_dir) / f"{exp_id}.json"
    write_json(target, {key: result[key] for key in ("cfg", "history", "summary")})
    return str(target)


def load_results(results_dir: str = "../results") -> list[dict]:
    results = [json.loads(path.read_text(encoding="utf-8")) for path in Path(results_dir).glob("*.json")]
    return sorted(results, key=lambda r: r["cfg"]["exp_id"])


def to_row(result: dict, eval_scores: dict | None = None, notes: str = "") -> dict:
    cfg, summary = result["cfg"], result["summary"]
    row = {**cfg, **summary}
    row["hidden"] = "-".join(map(str, cfg["hidden"]))
    row["loss"] = cfg["loss"].upper()
    row["optimizer"] = OPT_LABELS[cfg["optimizer"]]
    row["clip_norm"] = "none" if cfg["clip_norm"] is None else cfg["clip_norm"]
    row["diverged"] = "Y" if summary["diverged"] else "N"
    row["figure_file"] = f"figures/{cfg['exp_id']}.png"
    extra = [cfg.get("notes", ""), notes,
             f"momentum={cfg['momentum']}; betas={cfg.get('betas')}; eps={cfg.get('eps')}",
             f"train loss: fixed {summary.get('train_eval_size')} rows, eval mode",
             f"completed epochs={summary.get('completed_epochs')}",
             f"activation std after Linear={summary.get('activation_std')}"]
    if summary.get("divergence_reason"):
        extra.append(summary["divergence_reason"])
    if cfg.get("scheduler"):
        extra.append(f"scheduler={cfg['scheduler']}; {cfg.get('scheduler_kwargs', {})}")
    if summary.get("peak_mem_MB") is None:
        extra.append("GPU peak memory unavailable on CPU.")
    if cfg["loss"] == "mse":
        extra.append("MSE: mean over N*7 raw logits and one-hot targets; no 1/2 factor.")
    row["notes"] = "; ".join(item for item in extra if item)
    row["eval_acc"] = eval_scores["accuracy"] if eval_scores is not None else None
    row["eval_macro_f1"] = eval_scores["macro_f1"] if eval_scores is not None else None
    return json_safe(row)


def write_xlsx(rows: list[dict], template_path: str, out_path: str,
               seed_ids: list[str] | None = None, group_notes: dict | None = None,
               diagnostics: dict | None = None) -> None:
    ids = [row["exp_id"] for row in rows]
    if len(ids) != len(set(ids)):
        raise ValueError("Experiment IDs must be unique.")
    wb = openpyxl.load_workbook(template_path)
    if wb.sheetnames != ["Legend", "Experiments", "Seeds", "Summary"]:
        raise ValueError("Use the official four-sheet lab template.")
    ws = wb["Experiments"]
    headers = [cell.value for cell in ws[1]]
    original_last = ws.max_row
    last = max(original_last, len(rows) + 1)
    for r in range(2, last + 1):
        for c, name in enumerate(headers, 1):
            target = ws.cell(r, c)
            source = ws.cell(2, c)
            if r > original_last:
                target._style = copy(source._style)
                target.alignment = copy(source.alignment)
                ws.row_dimensions[r].height = ws.row_dimensions[2].height
            if name in FORMULA_COLUMNS:
                if r > original_last:
                    target.value = Translator(source.value, origin=source.coordinate).translate_formula(target.coordinate)
            else:
                target.value = rows[r - 2].get(name) if r - 2 < len(rows) else None
    # Keep template formulas and extend their lookup ranges when the table grows.
    if last > original_last:
        for sheet in wb:
            for row in sheet:
                for cell in row:
                    if cell.data_type == "f":
                        cell.value = re.sub(r"(\$[A-Z]+\$)" + str(original_last) + r"\b",
                                            lambda match: match.group(1) + str(last), cell.value)
        for validation in ws.data_validations.dataValidation:
            validation.sqref = str(validation.sqref).replace(str(original_last), str(last))
    if seed_ids is None:
        seed_ids = [row["exp_id"] for row in rows if row.get("group") == "baseline"][:3]
    if len(seed_ids) > 5 or any(exp_id not in ids for exp_id in seed_ids):
        raise ValueError("Seeds needs up to five existing baseline experiment IDs.")
    for r in range(2, 7):
        wb["Seeds"].cell(r, 1).value = seed_ids[r - 2] if r - 2 < len(seed_ids) else None
    for r in range(2, 12):
        group = wb["Summary"].cell(r, 1).value
        wb["Summary"].cell(r, 8).value = (group_notes or {}).get(group)
    if diagnostics is not None:
        # Keep diagnostic evidence within the existing Legend sheet.
        wb["Legend"]["A58"] = "Measured diagnostics (diagnostics.json)"
        for r, (key, value) in enumerate(diagnostics.items(), 59):
            wb["Legend"].cell(r, 1).value = key
            wb["Legend"].cell(r, 2).value = json.dumps(json_safe(value), ensure_ascii=True) if isinstance(value, (dict, list, tuple)) else value
    wb.calculation = CalcProperties(calcId=0, fullCalcOnLoad=True, forceFullCalc=True)
    target = Path(out_path)
    target.parent.mkdir(parents=True, exist_ok=True)
    wb.save(target)
