"""Regression checks for leakage, metrics, checkpoints, divergence and template formulas."""
from __future__ import annotations
import json
from pathlib import Path
import numpy as np
import pytest
import torch
from sklearn.metrics import f1_score
from data import apply_standardizer, fit_standardizer, iterate_batches, make_val_split
from model import MLP, EXPECTED_PARAMS, count_params, activation_stats
from optimizer import build_optimizer, build_scheduler, clip_gradients
from train import compute_loss, evaluate, macro_f1_from_confusion, run_experiment, final_eval, write_predictions, supports_bf16
from results_table import save_result, load_results, to_row, write_xlsx
from experiments import LabRunner, lr_trials, topic_trials, PREDICTIONS
from submission import verify_submission


@pytest.fixture(autouse=True)
def single_cpu_thread():
    torch.set_num_threads(1)


@pytest.fixture
def tiny_data():
    generator = torch.Generator().manual_seed(101)
    X = torch.randn(119, 54, generator=generator)
    y = torch.arange(119, dtype=torch.int64) % 7
    # Deliberately omit eval tensors: training must not require or access them.
    return {"X_tr": X[:91], "y_tr": y[:91], "X_val": X[91:], "y_val": y[91:]}


@pytest.mark.parametrize("hidden", list(EXPECTED_PARAMS))
def test_permitted_models(hidden):
    model = MLP(hidden, dropout=0.3)
    assert count_params(model) == EXPECTED_PARAMS[hidden]
    assert model(torch.randn(11, 54)).shape == (11, 7)
    assert not any(isinstance(layer, torch.nn.Softmax) for layer in model.modules())


def test_standardizer_does_not_leak_or_mutate():
    rng = np.random.default_rng(42)
    X = rng.normal(size=(140, 54)).astype(np.float32)
    X[:, 10:] = rng.integers(0, 2, size=(140, 44))
    X[:, 0] = 4
    y = (np.arange(140) % 7).astype(np.int64)
    original = X.copy()
    tr, yt, val, yv = make_val_split(X, y)
    mean, std = fit_standardizer(tr)
    val[:, 1] += 10_000
    scaled = apply_standardizer(tr, mean, std)
    assert np.array_equal(X, original)
    assert np.array_equal(scaled[:, 10:], tr[:, 10:])
    assert np.allclose(scaled[:, :10].mean(0), 0, atol=1e-6)
    assert np.allclose(scaled[:, 1:10].std(0), 1, atol=1e-6)
    assert np.all(scaled[:, 0] == 0)
    assert len(yt) == 112 and len(yv) == 28


def test_batches_keep_short_final_batch_and_seed(tiny_data):
    X, y = tiny_data["X_tr"], tiny_data["y_tr"]
    def shuffled():
        return list(iterate_batches(X, y, 20, torch.Generator().manual_seed(5)))
    first, second = shuffled(), shuffled()
    assert [len(x) for x, _ in first] == [20, 20, 20, 20, 11]
    assert torch.equal(torch.cat([x for x, _ in first]), torch.cat([x for x, _ in second]))
    assert torch.equal(torch.cat([x for x, _ in iterate_batches(X, y, 20, shuffle=False)]), X)


@pytest.mark.parametrize("loss_name", ["ce", "mse"])
def test_evaluation_sample_weighting_and_seven_classes(loss_name):
    torch.manual_seed(10)
    model = MLP(dropout=0.9)
    X, y = torch.randn(31, 54), torch.arange(31) % 3
    model.eval()
    with torch.no_grad():
        logits = model(X)
        expected = float(compute_loss(logits, y, loss_name))
        preds = logits.argmax(1).numpy()
    scores = evaluate(model, X, y, loss_name, batch_size=7)
    assert scores["loss"] == pytest.approx(expected, rel=1e-6)
    assert scores["macro_f1"] == pytest.approx(f1_score(y.numpy(), preds, labels=list(range(7)), average="macro", zero_division=0))
    assert not model.training


def test_zero_initialization_blocks_hidden_gradients():
    model = MLP(init="zeros")
    X, y = torch.randn(20, 54), torch.arange(20) % 7
    compute_loss(model(X), y, "ce").backward()
    assert model.net[0].weight.grad.count_nonzero() == 0
    assert model.net[2].weight.grad.count_nonzero() == 0
    assert model.net[-1].bias.grad.count_nonzero() > 0
    model.train()
    assert activation_stats(model, X) == [0, 0, 0]
    assert model.training


def test_mse_definition():
    logits = torch.tensor([[2., 0, 0, 0, 0, 0, 0]])
    assert float(compute_loss(logits, torch.tensor([0]), "mse")) == pytest.approx(1 / 7)


def test_preclip_norm_and_no_clip():
    p = torch.nn.Parameter(torch.zeros(2))
    p.grad = torch.tensor([3., 4.])
    assert clip_gradients([p], None) == 5
    assert torch.equal(p.grad, torch.tensor([3., 4.]))
    assert clip_gradients([p], 1.0) == 5
    assert float(p.grad.norm()) == pytest.approx(1, rel=1e-5)


def test_adam_and_adamw_equal_without_decay():
    a, b = torch.nn.Parameter(torch.ones(3)), torch.nn.Parameter(torch.ones(3))
    oa, ob = build_optimizer("adam", [a], 0.01), build_optimizer("adamw", [b], 0.01)
    for _ in range(4):
        a.grad = b.grad = torch.tensor([1., 2., 3.])
        oa.step()
        ob.step()
    assert torch.allclose(a, b)
    scheduler = build_scheduler(oa, "cosine", 10)
    assert scheduler.T_max == 10
    assert build_scheduler(ob, None, 1) is None


def test_training_deterministic_and_checkpoint_metrics(tiny_data):
    cfg = dict(exp_id="regression", lr=0.03, epochs=3, batch=19, train_eval_size=None)
    a, b = run_experiment(cfg, tiny_data), run_experiment(cfg, tiny_data)
    assert a["history"]["val_loss"] == b["history"]["val_loss"]
    assert a["summary"]["best_val_loss"] == min(a["history"]["val_loss"])
    model = MLP()
    model.load_state_dict(a["best_state"])
    score = evaluate(model, tiny_data["X_val"], tiny_data["y_val"])
    assert score["loss"] == pytest.approx(a["summary"]["best_val_loss"])
    assert score["macro_f1"] == pytest.approx(a["summary"]["val_macro_f1"])
    saved = a["best_state"]["net.0.weight"].clone()
    with torch.no_grad():
        model.net[0].weight.add_(100)
    assert torch.equal(saved, a["best_state"]["net.0.weight"])


def test_divergence_is_stopped_and_json_is_strict(tiny_data, tmp_path):
    result = run_experiment(dict(exp_id="diverged", lr=1e30, epochs=5, batch=20), tiny_data)
    assert result["summary"]["diverged"]
    assert result["summary"]["completed_epochs"] < 5
    path = save_result(result, tmp_path)
    text = Path(path).read_text()
    assert "NaN" not in text and "Infinity" not in text and "best_state" not in text
    assert load_results(tmp_path)[0]["cfg"]["exp_id"] == "diverged"


def test_prediction_export_and_official_id_alignment(tiny_data, tmp_path):
    result = run_experiment(dict(exp_id="predict", lr=0.03, epochs=1), tiny_data)
    data = {**tiny_data, "X_eval": tiny_data["X_val"], "eval_row_id": np.arange(28, dtype=np.int64)[::-1]}
    target = tmp_path / "pred.csv"
    final_eval(result["cfg"], result, data, str(target))
    assert target.read_text().splitlines()[0] == "row_id,pred"
    values = np.loadtxt(target, delimiter=",", skiprows=1, dtype=np.int64)
    assert np.array_equal(values[:, 0], data["eval_row_id"])
    with pytest.raises(ValueError):
        write_predictions(np.array([1, 1]), np.array([0, 1]), str(target))
    with pytest.raises(ValueError):
        final_eval({**result["cfg"], "seed": 2}, result, data, str(target))


def test_cache_reuse_and_stale_configuration_rejected(tiny_data, tmp_path):
    out = tmp_path / "submission_test"
    cfg = dict(exp_id="cached", lr=0.03, epochs=1)
    runner = LabRunner(tiny_data, tmp_path, out)
    first = runner.run(cfg)
    second = LabRunner(tiny_data, tmp_path, out).run(cfg)
    assert first["summary"]["val_acc"] == second["summary"]["val_acc"]
    assert torch.equal(first["best_state"]["net.0.weight"], second["best_state"]["net.0.weight"])
    with pytest.raises(RuntimeError, match="different code"):
        runner.run({**cfg, "lr": 0.04})
    changed = {**tiny_data, "X_tr": tiny_data["X_tr"] + 0.1}
    old_environment = (out / "environment.json").read_bytes()
    changed_runner = LabRunner(changed, tmp_path, out)
    assert changed_runner.out != out
    changed_runner.run(cfg)
    assert (out / "environment.json").read_bytes() == old_environment


@pytest.mark.parametrize("frozen", [False, True])
def test_submitted_logs_without_checkpoints_start_an_isolated_run(tiny_data, tmp_path, frozen):
    out = tmp_path / "submission_test"
    cfg = dict(exp_id="cached", lr=0.03, epochs=1)
    runner = LabRunner(tiny_data, tmp_path, out)
    runner.run(cfg)
    record = out / "results/cached.json"
    original_record = record.read_bytes()
    original_environment = (out / "environment.json").read_bytes()
    (runner.checkpoints / "cached.pt").unlink()
    if frozen:
        (out / "selection.json").write_text("{}", encoding="utf-8")
    replay = LabRunner(tiny_data, tmp_path, out)
    assert replay.out != out
    assert replay.out.name == out.name
    assert ".lab_reruns" in replay.out.parts
    assert not (replay.out / "selection.json").exists()
    replay.run(cfg)
    assert (replay.checkpoints / "cached.pt").exists()
    assert record.read_bytes() == original_record
    assert (out / "environment.json").read_bytes() == original_environment


def test_completed_submission_is_preserved_even_with_checkpoints(tiny_data, tmp_path):
    out = tmp_path / "submission_test"
    runner = LabRunner(tiny_data, tmp_path, out)
    runner.run(dict(exp_id="cached", lr=0.03, epochs=1))
    (out / "selection.json").write_text("{}", encoding="utf-8")
    assert LabRunner(tiny_data, tmp_path, out).out != out


def test_template_preservation_and_more_than_sixty_rows(tiny_data, tmp_path):
    import openpyxl
    result = run_experiment(dict(exp_id="base-s1", lr=0.03, epochs=1), tiny_data)
    rows = [{**to_row(result), "exp_id": f"row-{i}"} for i in range(65)]
    target = tmp_path / "table.xlsx"
    template = Path(__file__).with_name("experiment_table_template.xlsx")
    write_xlsx(rows, str(template), str(target), seed_ids=["row-0", "row-1", "row-2"])
    wb = openpyxl.load_workbook(target)
    source = openpyxl.load_workbook(template)
    assert wb.sheetnames == source.sheetnames
    assert [c.value for c in wb["Experiments"][1]] == [c.value for c in source["Experiments"][1]]
    assert wb["Experiments"]["AD2"].value == source["Experiments"]["AD2"].value
    assert wb["Experiments"]["AD66"].value == '=IF(P66="","",P66-LN(7))'
    assert "$66" in wb["Seeds"]["B2"].value
    assert wb["Seeds"]["A2"].value == "row-0"
    assert wb["Experiments"]["D2"].value == "CE"
    assert wb["Experiments"]["E2"].value == "SGD+momentum"
    assert wb.calculation.fullCalcOnLoad


def test_mixed_precision_requires_cuda(tiny_data):
    with pytest.raises(ValueError, match="CUDA"):
        run_experiment(dict(lr=0.03, epochs=1, precision="fp16"), tiny_data)


def test_bf16_emulation_is_not_native_support(monkeypatch):
    monkeypatch.setattr(torch.cuda, "is_available", lambda: True)
    monkeypatch.setattr(torch.cuda, "is_bf16_supported", lambda: True)
    monkeypatch.setattr(torch.cuda, "get_device_capability", lambda device=None: (7, 5))
    assert not supports_bf16()
    monkeypatch.setattr(torch.cuda, "get_device_capability", lambda device=None: (8, 6))
    assert supports_bf16()


def test_topic_predictions_do_not_inherit_baseline_prediction(tiny_data):
    baseline = run_experiment(dict(lr=0.03, epochs=1, prediction="Baseline prediction"), tiny_data)
    for cfg in topic_trials(baseline, torch.device("cpu")):
        assert cfg["prediction"] == PREDICTIONS[cfg["group"]]


def test_smoke_outputs_cannot_be_submitted(tmp_path):
    (tmp_path / "environment.json").write_text(json.dumps({"smoke_test": True}))
    with pytest.raises(ValueError, match="Smoke-test"):
        verify_submission(tmp_path)


def test_notebook_syntax_and_embedded_modules():
    import base64
    import io
    import zipfile
    nb = json.loads(Path(__file__).with_name("lab.ipynb").read_text(encoding="utf-8"))
    for c in nb["cells"]:
        if isinstance(c["source"], list):
            c["source"] = "".join(c["source"])
    for c in nb["cells"]:
        if c["cell_type"] == "code":
            compile(c["source"], "notebook-cell", "exec")
    environment = nb["cells"][2]["source"]
    assignment = next(line for line in environment.splitlines() if line.startswith("BUNDLE = "))
    namespace = {}
    exec(assignment, namespace)
    with zipfile.ZipFile(io.BytesIO(base64.b64decode(namespace["BUNDLE"]))) as archive:
        for name in ("data.py", "model.py", "optimizer.py", "train.py", "experiments.py", "submission.py"):
            assert archive.read(name) == Path(__file__).with_name(name).read_bytes()
