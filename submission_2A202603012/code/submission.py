"""Freeze validation selection, run the official evaluator and package checked outputs."""
from __future__ import annotations
import json
from pathlib import Path
import shutil
import subprocess
import sys
import zipfile
import numpy as np
import matplotlib.pyplot as plt
import openpyxl
from experiments import best_by_validation
from train import final_eval
from results_table import write_json, save_result, to_row, write_xlsx
from report import observation, write_report


def official_evaluate(root, pred_path, result_path):
    root = Path(root).resolve()
    subprocess.run([sys.executable, "-X", "utf8", str(root / "scripts/evaluate.py"),
                    "--data", str(root / "data/covtype.csv.gz"),
                    "--meta", str(root / "data/split_metadata.csv"),
                    "--pred", str(Path(pred_path).resolve()),
                    "--out", str(Path(result_path).resolve())], cwd=root, check=True)
    return json.loads(Path(result_path).read_text(encoding="utf-8"))


def finalize(runner, baseline, baseline_seeds, seed_stats, diagnostics, student_id):
    results = list(runner.results.values())
    # Compare seed 1 configurations; selecting the best seed would bias the comparison.
    final = best_by_validation([r for r in results if r["cfg"]["seed"] == 1])
    selection = {
        "baseline": baseline["cfg"]["exp_id"], "final": final["cfg"]["exp_id"],
        "baseline_fingerprint": baseline["summary"]["fingerprint"],
        "final_fingerprint": final["summary"]["fingerprint"],
        "candidate_ids": sorted(runner.results),
        "rule": "highest validation macro-F1 at each minimum-val-loss checkpoint; seed 1 only",
    }
    selection_path = runner.out / "selection.json"
    if selection_path.exists() and json.loads(selection_path.read_text(encoding="utf-8")) != selection:
        raise RuntimeError("Final selection was already frozen. Do not retune after viewing eval.")
    write_json(selection_path, selection)
    final_eval(final["cfg"], final, runner.data, str(runner.out / "predictions_eval.csv"))
    final_scores = official_evaluate(runner.root, runner.out / "predictions_eval.csv", runner.out / "eval_result.json")
    if baseline["cfg"]["exp_id"] == final["cfg"]["exp_id"]:
        base_scores = final_scores
        write_json(runner.out / "baseline_eval_result.json", base_scores)
    else:
        final_eval(baseline["cfg"], baseline, runner.data, str(runner.out / "baseline_predictions.csv"))
        base_scores = official_evaluate(runner.root, runner.out / "baseline_predictions.csv",
                                        runner.out / "baseline_eval_result.json")
    scores = {baseline["cfg"]["exp_id"]: base_scores, final["cfg"]["exp_id"]: final_scores}
    noise = seed_stats["val_macro_f1"]["noise_2sigma"]
    group_notes = {}
    for result in results:
        reference = runner.results.get(result["cfg"].get("reference"), baseline)
        if reference["summary"]["diverged"]:
            reference = baseline
        text = observation(result, reference, noise)
        result["summary"]["observation"] = text
        save_result(result, runner.out / "results")
        group_notes[result["cfg"]["group"]] = text
    if not any(r["cfg"]["precision"] == "bf16" for r in results):
        group_notes["amp"] = group_notes.get("amp", "No CUDA runtime; mixed precision skipped.") + " BF16 not supported by this runtime."
    ordered = [baseline] + sorted([r for r in results if r is not baseline], key=lambda r: r["cfg"]["exp_id"])
    rows = [to_row(r, scores.get(r["cfg"]["exp_id"]),
                   notes=f"Prediction: {r['cfg'].get('prediction', '')}; Observation: {r['summary']['observation']}")
            for r in ordered]
    template = Path(__file__).with_name("experiment_table_template.xlsx")
    if not template.exists():
        template = runner.root / "templates/experiment_table_template.xlsx"
    write_xlsx(rows, str(template), str(runner.out / "experiments.xlsx"),
               seed_ids=[r["cfg"]["exp_id"] for r in baseline_seeds],
               group_notes=group_notes, diagnostics=diagnostics)
    runner.comparisons(baseline)
    cm = np.asarray(final_scores["confusion_matrix"])
    fig, ax = plt.subplots(figsize=(6, 5))
    im = ax.imshow(cm, cmap="Blues")
    ax.set(xlabel="Predicted class", ylabel="True class", title="Official eval confusion matrix",
           xticks=range(7), yticks=range(7))
    for i in range(7):
        for j in range(7):
            ax.text(j, i, str(cm[i, j]), ha="center", va="center", fontsize=8,
                    color="white" if cm[i, j] > cm.max() / 2 else "black")
    fig.colorbar(im, ax=ax)
    fig.tight_layout()
    fig.savefig(runner.out / "figures/eval_confusion.png", dpi=160)
    plt.close(fig)
    report = runner.out / "REPORT.md"
    report_template = runner.root / "templates/REPORT_TEMPLATE.md"
    if report.exists() and not report.read_text(encoding="utf-8").startswith("# Lab Day 1 report -"):
        if not report_template.exists() or report.read_bytes() != report_template.read_bytes():
            report = runner.out / "REPORT.generated.md"
    write_report(report, ordered, baseline, final, seed_stats, scores, diagnostics,
                 runner.environment, student_id)
    print("Selected final configuration:", selection["final"])
    print("Report draft:", report)
    print("Open experiments.xlsx in Excel and save to refresh its formula caches.")
    return {"selection": selection, "scores": scores, "final": final, "baseline": baseline}


def verify_submission(out_dir, require_executed_notebook=True):
    out = Path(out_dir)
    environment_path = out / "environment.json"
    if environment_path.exists() and json.loads(environment_path.read_text(encoding="utf-8")).get("smoke_test"):
        raise ValueError("Smoke-test outputs cannot be packaged as a lab submission.")
    required = ["REPORT.md", "experiments.xlsx", "predictions_eval.csv", "eval_result.json", "code/lab.ipynb"]
    missing = [name for name in required if not (out / name).is_file()]
    if missing:
        raise ValueError(f"Missing submission files: {missing}")
    report = (out / "REPORT.md").read_text(encoding="utf-8")
    if "<MSSV>" in report or "___" in report or "<H? t?n>" in report:
        raise ValueError("REPORT.md still contains template placeholders.")
    notebook = json.loads((out / "code/lab.ipynb").read_text(encoding="utf-8"))
    cells = [c for c in notebook["cells"] if c["cell_type"] == "code"]
    if any(o.get("output_type") == "error" for c in cells for o in c.get("outputs", [])):
        raise ValueError("The notebook contains an error output.")
    if require_executed_notebook and any(c.get("execution_count") is None for c in cells):
        raise ValueError("Save the executed notebook with outputs before packaging.")
    wb = openpyxl.load_workbook(out / "experiments.xlsx", data_only=False)
    if wb.sheetnames != ["Legend", "Experiments", "Seeds", "Summary"]:
        raise ValueError("Workbook sheet names do not match the template.")
    ws = wb["Experiments"]
    headers = {c.value: c.column for c in ws[1]}
    ids = []
    for row in range(2, ws.max_row + 1):
        exp_id = ws.cell(row, headers["exp_id"]).value
        if not exp_id:
            continue
        ids.append(exp_id)
        figure = ws.cell(row, headers["figure_file"]).value
        if figure != f"figures/{exp_id}.png" or not (out / figure).is_file():
            raise ValueError(f"Missing or mismatched figure for {exp_id}.")
        record = out / "results" / f"{exp_id}.json"
        if not record.exists():
            raise ValueError(f"Missing result record for {exp_id}.")
    if not ids or len(ids) != len(set(ids)):
        raise ValueError("Workbook must contain unique experiment rows.")
    score = json.loads((out / "eval_result.json").read_text(encoding="utf-8"))
    selected = json.loads((out / "selection.json").read_text(encoding="utf-8"))
    for row in range(2, ws.max_row + 1):
        exp_id = ws.cell(row, headers["exp_id"]).value
        values = [ws.cell(row, headers[k]).value for k in ("eval_acc", "eval_macro_f1")]
        if exp_id not in {selected["baseline"], selected["final"]} and any(v is not None for v in values):
            raise ValueError("Eval scores appear in a non-final experiment row.")
        if exp_id == selected["final"]:
            if values != [score["accuracy"], score["macro_f1"]]:
                if not all(v is not None and abs(v - expected) < 1e-12
                           for v, expected in zip(values, [score["accuracy"], score["macro_f1"]])):
                    raise ValueError("Final workbook scores disagree with the official evaluator.")
    for file in (out / "code").glob("*.py"):
        compile(file.read_text(encoding="utf-8"), str(file), "exec")
    return {"experiments": len(ids), "n_eval": score["n_eval"], "final": selected["final"]}


def package_submission(out_dir, zip_path=None):
    out = Path(out_dir).resolve()
    if not out.name.startswith("submission_"):
        raise ValueError("The submission folder must be named submission_<student_id>.")
    verify_submission(out)
    target = Path(zip_path or out.parent / f"{out.name}.zip").resolve()
    if out == target or out in target.parents:
        raise ValueError("The ZIP must stay outside the submission folder.")
    top_files = {"REPORT.md", "experiments.xlsx", "predictions_eval.csv", "eval_result.json",
                 "baseline_eval_result.json", "selection.json", "diagnostics.json", "environment.json", "seed_statistics.json"}
    excluded = {"colab_setup.ipynb", "colab_bootstrap.py"}
    included = []
    for file in out.rglob("*"):
        relative = file.relative_to(out)
        if not file.is_file() or any(part in {"__pycache__", ".ipynb_checkpoints"} for part in relative.parts):
            continue
        if file.name in excluded or file.suffix in {".pt", ".pth", ".ckpt", ".npz", ".pyc"}:
            continue
        if len(relative.parts) == 1 and file.name in top_files:
            included.append(file)
        elif relative.parts[0] == "code" and file.suffix in {".py", ".ipynb", ".txt", ".xlsx", ".md"}:
            included.append(file)
        elif relative.parts[0] == "figures" and file.suffix == ".png":
            included.append(file)
        elif relative.parts[0] == "results" and file.suffix == ".json":
            included.append(file)
    target.parent.mkdir(parents=True, exist_ok=True)
    with zipfile.ZipFile(target, "w", zipfile.ZIP_DEFLATED) as archive:
        for file in included:
            archive.write(file, f"{out.name}/{file.relative_to(out).as_posix()}")
    print("Submission ZIP:", target)
    return target
